"""Read text out of existing PDF/DOCX files (as opposed to pptx_tools.py /
docx_tools.py / xlsx_tools.py, which create and edit new ones) so they can
be summarized or questioned — 'what does this PDF say', 'summarize this
report'. Reads local file paths; combine with an interface's file-download
step (e.g. a Telegram document) to point this at a received file."""
from __future__ import annotations

from docx import Document as WordDocument
from pypdf import PdfReader

from tools.base import Tool


class ReadPdfTool(Tool):
    requires_network = False
    name = "read_pdf"
    description = "Extract text from a PDF file by its local path, for summarizing or answering questions about it."
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "max_chars": {"type": "integer", "description": "Default 8000."},
        },
        "required": ["path"],
    }

    def run(self, path: str, max_chars: int = 8000) -> str:
        reader = PdfReader(path)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            return "No extractable text found (the PDF may be scanned images without OCR)."
        return text[:max_chars]


class ReadWordDocumentTool(Tool):
    requires_network = False
    name = "read_word_document"
    description = "Extract text from an existing Word document (.docx) by its local path."
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "max_chars": {"type": "integer", "description": "Default 8000."},
        },
        "required": ["path"],
    }

    def run(self, path: str, max_chars: int = 8000) -> str:
        doc = WordDocument(path)
        # doc.paragraphs only walks top-level body paragraphs — it skips
        # table content entirely, so a document containing a table (created
        # by this project's own add_table_to_word_document, or received/
        # opened from anywhere else) used to have that data silently
        # omitted with no indication anything was skipped.
        parts = [p.text for p in doc.paragraphs]
        for i, table in enumerate(doc.tables, start=1):
            parts.append(f"\n[Table {i}]")
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
        return text[:max_chars] if text.strip() else "No text found in this document."
