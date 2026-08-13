import os

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from core.store import Store
from tools.pptx_tools import (
    AddImageToSlideTool,
    AddSlideTool,
    CreatePresentationTool,
    DeleteSlideTool,
    EditSlideTool,
    ListSlidesTool,
    SetSlideNotesTool,
)


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return Store(db_path=str(tmp_path / "test.db"))


def test_create_presentation_derives_name_and_saves_file(store):
    result = CreatePresentationTool(store).run(
        title="My Deck", slides=[{"title": "Intro", "bullets": ["a", "b"]}]
    )
    assert "my-deck" in result

    doc = store.get_document("my-deck")
    assert doc is not None
    assert os.path.exists(doc.path)

    prs = Presentation(doc.path)
    assert len(prs.slides) == 2  # title slide + one content slide


def test_add_slide_appends(store):
    CreatePresentationTool(store).run(title="My Deck")
    AddSlideTool(store).run(document_name="my-deck", title="New Slide", bullets=["x"])

    doc = store.get_document("my-deck")
    prs = Presentation(doc.path)
    assert len(prs.slides) == 2
    assert prs.slides[1].shapes.title.text == "New Slide"


def test_edit_slide_changes_title(store):
    CreatePresentationTool(store).run(title="My Deck", slides=[{"title": "Old", "bullets": ["a"]}])
    EditSlideTool(store).run(document_name="my-deck", slide_index=2, title="New Title")

    doc = store.get_document("my-deck")
    prs = Presentation(doc.path)
    assert prs.slides[1].shapes.title.text == "New Title"


def test_delete_slide_removes_it(store):
    CreatePresentationTool(store).run(
        title="My Deck",
        slides=[{"title": "First", "bullets": ["a"]}, {"title": "Second", "bullets": ["b"]}],
    )
    DeleteSlideTool(store).run(document_name="my-deck", slide_index=2)

    doc = store.get_document("my-deck")
    prs = Presentation(doc.path)
    assert len(prs.slides) == 2
    assert prs.slides[1].shapes.title.text == "Second"


def test_list_slides(store):
    CreatePresentationTool(store).run(title="My Deck", slides=[{"title": "Only slide"}])
    result = ListSlidesTool(store).run(document_name="my-deck")
    assert "My Deck" in result
    assert "Only slide" in result


def test_operations_on_unknown_document_return_friendly_error(store):
    result = AddSlideTool(store).run(document_name="does-not-exist", title="X")
    assert "No presentation" in result


def test_create_presentation_warns_on_cross_kind_name_collision(store):
    store.register_document("xlsx", "budget", "/outputs/budget.xlsx")
    result = CreatePresentationTool(store).run(title="Budget", document_name="budget")
    assert "replaced an existing xlsx document" in result


def test_add_image_to_slide(store, tmp_path):
    CreatePresentationTool(store).run(title="My Deck", slides=[{"title": "Intro"}])
    image_path = tmp_path / "photo.png"
    Image.new("RGB", (40, 30), color="blue").save(image_path)

    result = AddImageToSlideTool(store).run(
        document_name="my-deck", slide_index=2, image_path=str(image_path)
    )

    assert "Added image" in result
    doc = store.get_document("my-deck")
    prs = Presentation(doc.path)
    slide = prs.slides[1]
    picture_shapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(picture_shapes) == 1


def test_add_image_to_slide_missing_file(store):
    CreatePresentationTool(store).run(title="My Deck", slides=[{"title": "Intro"}])
    result = AddImageToSlideTool(store).run(
        document_name="my-deck", slide_index=2, image_path="nope.png"
    )
    assert "is not a file" in result


def test_add_image_to_slide_invalid_index(store, tmp_path):
    CreatePresentationTool(store).run(title="My Deck", slides=[{"title": "Intro"}])
    image_path = tmp_path / "photo.png"
    Image.new("RGB", (10, 10)).save(image_path)

    result = AddImageToSlideTool(store).run(document_name="my-deck", slide_index=99, image_path=str(image_path))
    assert "only has" in result


def test_set_slide_notes(store):
    CreatePresentationTool(store).run(title="My Deck", slides=[{"title": "Intro"}])
    result = SetSlideNotesTool(store).run(document_name="my-deck", slide_index=2, notes="Remember to smile.")

    assert "Set speaker notes" in result
    doc = store.get_document("my-deck")
    prs = Presentation(doc.path)
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Remember to smile."


def test_set_slide_notes_unknown_document(store):
    result = SetSlideNotesTool(store).run(document_name="ghost", slide_index=1, notes="x")
    assert "No presentation" in result
