import pytest
from openpyxl import load_workbook

from core.store import Store
from tools.xlsx_tools import AddSpreadsheetChartTool, CreateSpreadsheetTool, SetSpreadsheetFormulaTool


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return Store(db_path=str(tmp_path / "test.db"))


def test_set_formula(store):
    CreateSpreadsheetTool(store).run(title="Budget", headers=["Item", "Price"], rows=[["Bread", 2.5]])
    SetSpreadsheetFormulaTool(store).run(document_name="budget", cell="B3", formula="=SUM(B2:B2)")

    doc = store.get_document("budget")
    wb = load_workbook(doc.path)
    assert wb.active["B3"].value == "=SUM(B2:B2)"


def test_set_formula_unknown_document(store):
    result = SetSpreadsheetFormulaTool(store).run(document_name="ghost", cell="A1", formula="=1+1")
    assert "No spreadsheet" in result


def test_add_chart(store):
    CreateSpreadsheetTool(store).run(
        title="Budget", headers=["Item", "Price"], rows=[["Bread", 2.5], ["Milk", 1.2]]
    )
    result = AddSpreadsheetChartTool(store).run(
        document_name="budget", chart_type="bar", data_range="B2:B3", category_range="A2:A3"
    )
    assert "bar chart" in result

    doc = store.get_document("budget")
    wb = load_workbook(doc.path)
    assert len(wb.active._charts) == 1


def test_add_chart_unknown_document(store):
    result = AddSpreadsheetChartTool(store).run(
        document_name="ghost", chart_type="pie", data_range="A1:A2", category_range="B1:B2"
    )
    assert "No spreadsheet" in result
