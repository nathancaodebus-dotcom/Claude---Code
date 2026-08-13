from docx import Document as WordDocument
from pypdf import PdfWriter

from tools.document_reader_tools import ReadPdfTool, ReadWordDocumentTool


def test_read_word_document(tmp_path):
    path = tmp_path / "test.docx"
    doc = WordDocument()
    doc.add_paragraph("Hello from a test document.")
    doc.save(path)

    result = ReadWordDocumentTool().run(path=str(path))
    assert "Hello from a test document." in result


def test_read_word_document_includes_table_content(tmp_path):
    """Regression test: doc.paragraphs only walks top-level body
    paragraphs and skips tables entirely — any document with a table
    (created by this project's own add_table_to_word_document, or received
    from anywhere else) used to have that data silently omitted."""
    path = tmp_path / "test.docx"
    doc = WordDocument()
    doc.add_paragraph("Intro text.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Bread"
    table.cell(1, 1).text = "2.50"
    doc.save(path)

    result = ReadWordDocumentTool().run(path=str(path))

    assert "Intro text." in result
    assert "Name" in result and "Price" in result
    assert "Bread" in result and "2.50" in result


def test_read_pdf_with_no_extractable_text(tmp_path):
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as f:
        writer.write(f)

    result = ReadPdfTool().run(path=str(path))
    assert "No extractable text" in result
