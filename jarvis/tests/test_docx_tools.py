import os

import pytest
from docx import Document as WordDocument

from core.store import Store
from tools.docx_tools import AppendToWordDocumentTool, CreateWordDocumentTool, ListWordDocumentsTool


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return Store(db_path=str(tmp_path / "test.db"))


def test_create_word_document(store):
    result = CreateWordDocumentTool(store).run(title="My Report", paragraphs=["first"])
    assert "my-report" in result

    doc = store.get_document("my-report")
    assert os.path.exists(doc.path)
    word_doc = WordDocument(doc.path)
    assert word_doc.paragraphs[0].text == "My Report"
    assert word_doc.paragraphs[1].text == "first"


def test_append_to_word_document(store):
    CreateWordDocumentTool(store).run(title="My Report")
    AppendToWordDocumentTool(store).run(document_name="my-report", text="second paragraph")

    doc = store.get_document("my-report")
    word_doc = WordDocument(doc.path)
    assert word_doc.paragraphs[-1].text == "second paragraph"


def test_append_unknown_document_returns_friendly_error(store):
    result = AppendToWordDocumentTool(store).run(document_name="ghost", text="x")
    assert "No Word document" in result


def test_list_word_documents(store):
    CreateWordDocumentTool(store).run(title="Report One")
    CreateWordDocumentTool(store).run(title="Report Two")

    result = ListWordDocumentsTool(store).run()
    assert "report-one" in result
    assert "report-two" in result
