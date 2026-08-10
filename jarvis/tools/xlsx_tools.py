"""Create and edit Excel spreadsheets by voice or text."""
from __future__ import annotations

from openpyxl import Workbook, load_workbook

from core.attachments import push as push_attachment
from core.store import Store
from tools.base import Tool
from tools.document_utils import resolve_path, slugify


class CreateSpreadsheetTool(Tool):
    name = "create_spreadsheet"
    description = (
        "Create a new Excel spreadsheet with a header row and data rows. Returns the "
        "document name to use in later add_spreadsheet_row calls."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "title": {"type": "string"},
            "headers": {"type": "array", "items": {"type": "string"}},
            "rows": {
                "type": "array",
                "items": {"type": "array", "items": {"type": ["string", "number"]}},
                "description": "Each item is one row, a list of cell values matching the headers.",
            },
            "document_name": {"type": "string", "description": "Derived from the title if omitted."},
        },
        "required": ["title", "headers"],
    }

    def __init__(self, store: Store):
        self._store = store

    def run(self, title: str, headers: list[str], rows: list[list] | None = None, document_name: str | None = None) -> str:
        document_name = document_name or slugify(title)
        path = resolve_path(document_name, "xlsx")

        wb = Workbook()
        sheet = wb.active
        sheet.title = title[:31]  # Excel sheet name limit
        sheet.append(headers)
        for row in rows or []:
            sheet.append(row)
        wb.save(path)

        self._store.register_document("xlsx", document_name, str(path))
        push_attachment(str(path))
        return f"Created spreadsheet '{document_name}' with {len(rows or [])} data rows, saved to {path}."


class AddSpreadsheetRowTool(Tool):
    name = "add_spreadsheet_row"
    description = "Append a row of values to an existing spreadsheet."
    input_schema = {
        "type": "object",
        "properties": {
            "document_name": {"type": "string"},
            "row": {"type": "array", "items": {"type": ["string", "number"]}},
        },
        "required": ["document_name", "row"],
    }

    def __init__(self, store: Store):
        self._store = store

    def run(self, document_name: str, row: list) -> str:
        doc = self._store.get_document(document_name)
        if not doc or doc.kind != "xlsx":
            return f"No spreadsheet named '{document_name}'. Use create_spreadsheet first."

        wb = load_workbook(doc.path)
        wb.active.append(row)
        wb.save(doc.path)

        self._store.touch_document(document_name)
        push_attachment(doc.path)
        return f"Added row to '{document_name}'."


class ListSpreadsheetsTool(Tool):
    name = "list_spreadsheets"
    description = "List all spreadsheets Jarvis has created."
    input_schema = {"type": "object", "properties": {}}

    def __init__(self, store: Store):
        self._store = store

    def run(self) -> str:
        docs = self._store.list_documents(kind="xlsx")
        if not docs:
            return "No spreadsheets yet."
        return "\n".join(f"- {d.name} ({d.path})" for d in docs)
