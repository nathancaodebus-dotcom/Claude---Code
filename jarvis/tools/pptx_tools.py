"""Create and edit PowerPoint presentations by voice or text — 'create me a
PowerPoint about X' then 'add a slide about Y' / 'change slide 2's title'.

Presentations are tracked by name in the Store so a follow-up command can
find the same file again without the user repeating the full path. If the
user doesn't give a name, one is derived from the title and returned in the
tool result so the model can refer back to it in later turns.
"""
from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt

from core.attachments import push as push_attachment
from core.store import Store
from tools.base import Tool
from tools.document_utils import resolve_path, slugify

_TITLE_AND_CONTENT_LAYOUT = 1
_TITLE_ONLY_LAYOUT = 5


def _add_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    layout = prs.slide_layouts[_TITLE_AND_CONTENT_LAYOUT if bullets else _TITLE_ONLY_LAYOUT]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title

    if bullets:
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            paragraph = body.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0


class CreatePresentationTool(Tool):
    name = "create_presentation"
    description = (
        "Create a new PowerPoint presentation with a title slide and a list of content "
        "slides (each with a title and bullet points). Returns the document name to use "
        "in later add_slide/edit_slide/delete_slide calls."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "title": {"type": "string", "description": "Presentation title (used for the title slide)."},
            "subtitle": {"type": "string", "description": "Optional subtitle for the title slide."},
            "slides": {
                "type": "array",
                "description": "Content slides to add after the title slide.",
                "items": {
                    "type": "object",
                    "properties": {
                        "title": {"type": "string"},
                        "bullets": {"type": "array", "items": {"type": "string"}},
                    },
                    "required": ["title"],
                },
            },
            "document_name": {
                "type": "string",
                "description": "Optional stable name for this document. Derived from the title if omitted.",
            },
        },
        "required": ["title"],
    }

    def __init__(self, store: Store):
        self._store = store

    def run(self, title: str, subtitle: str = "", slides: list[dict] | None = None, document_name: str | None = None) -> str:
        document_name = document_name or slugify(title)
        path = resolve_path(document_name, "pptx")

        prs = Presentation()
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = title
        if subtitle:
            title_slide.placeholders[1].text = subtitle

        for slide in slides or []:
            _add_slide(prs, slide["title"], slide.get("bullets", []))

        prs.save(path)
        self._store.register_document("pptx", document_name, str(path))
        push_attachment(str(path))
        return (
            f"Created presentation '{document_name}' with {1 + len(slides or [])} slides, "
            f"saved to {path}."
        )


class AddSlideTool(Tool):
    name = "add_slide"
    description = "Add a new content slide (title + bullet points) to an existing presentation."
    input_schema = {
        "type": "object",
        "properties": {
            "document_name": {"type": "string"},
            "title": {"type": "string"},
            "bullets": {"type": "array", "items": {"type": "string"}},
        },
        "required": ["document_name", "title"],
    }

    def __init__(self, store: Store):
        self._store = store

    def run(self, document_name: str, title: str, bullets: list[str] | None = None) -> str:
        doc = self._store.get_document(document_name)
        if not doc or doc.kind != "pptx":
            return f"No presentation named '{document_name}'. Use create_presentation first."

        prs = Presentation(doc.path)
        _add_slide(prs, title, bullets or [])
        prs.save(doc.path)
        self._store.touch_document(document_name)
        push_attachment(doc.path)
        return f"Added slide '{title}' to '{document_name}' (now {len(prs.slides)} slides)."


class EditSlideTool(Tool):
    name = "edit_slide"
    description = (
        "Edit an existing slide's title and/or bullet points by its 1-based index "
        "(slide 1 is the title slide). Use list_slides first if unsure of the index."
    )
    input_schema = {
        "type": "object",
        "properties": {
            "document_name": {"type": "string"},
            "slide_index": {"type": "integer", "description": "1-based slide number."},
            "title": {"type": "string", "description": "New title. Omit to leave unchanged."},
            "bullets": {
                "type": "array",
                "items": {"type": "string"},
                "description": "New bullet points, replacing the existing ones. Omit to leave unchanged.",
            },
        },
        "required": ["document_name", "slide_index"],
    }

    def __init__(self, store: Store):
        self._store = store

    def run(self, document_name: str, slide_index: int, title: str | None = None, bullets: list[str] | None = None) -> str:
        doc = self._store.get_document(document_name)
        if not doc or doc.kind != "pptx":
            return f"No presentation named '{document_name}'."

        prs = Presentation(doc.path)
        slides = list(prs.slides)
        if not (1 <= slide_index <= len(slides)):
            return f"'{document_name}' only has {len(slides)} slides."

        slide = slides[slide_index - 1]
        if title is not None and slide.shapes.title is not None:
            slide.shapes.title.text = title

        if bullets is not None and len(slide.placeholders) > 1:
            body = slide.placeholders[1].text_frame
            body.clear()
            body.text = bullets[0] if bullets else ""
            for bullet in bullets[1:]:
                paragraph = body.add_paragraph()
                paragraph.text = bullet

        prs.save(doc.path)
        self._store.touch_document(document_name)
        push_attachment(doc.path)
        return f"Updated slide {slide_index} in '{document_name}'."


class DeleteSlideTool(Tool):
    name = "delete_slide"
    description = "Delete a slide from a presentation by its 1-based index."
    input_schema = {
        "type": "object",
        "properties": {
            "document_name": {"type": "string"},
            "slide_index": {"type": "integer"},
        },
        "required": ["document_name", "slide_index"],
    }

    def __init__(self, store: Store):
        self._store = store

    def run(self, document_name: str, slide_index: int) -> str:
        doc = self._store.get_document(document_name)
        if not doc or doc.kind != "pptx":
            return f"No presentation named '{document_name}'."

        prs = Presentation(doc.path)
        slides = list(prs.slides)
        if not (1 <= slide_index <= len(slides)):
            return f"'{document_name}' only has {len(slides)} slides."

        xml_slides = prs.slides._sldIdLst
        slide_id = list(xml_slides)[slide_index - 1]
        xml_slides.remove(slide_id)

        prs.save(doc.path)
        self._store.touch_document(document_name)
        push_attachment(doc.path)
        return f"Deleted slide {slide_index} from '{document_name}'."


class ListSlidesTool(Tool):
    name = "list_slides"
    description = "List the slides in a presentation (index and title)."
    input_schema = {
        "type": "object",
        "properties": {"document_name": {"type": "string"}},
        "required": ["document_name"],
    }

    def __init__(self, store: Store):
        self._store = store

    def run(self, document_name: str) -> str:
        doc = self._store.get_document(document_name)
        if not doc or doc.kind != "pptx":
            return f"No presentation named '{document_name}'."

        prs = Presentation(doc.path)
        lines = []
        for i, slide in enumerate(prs.slides, start=1):
            title = slide.shapes.title.text if slide.shapes.title else "(no title)"
            lines.append(f"{i}. {title}")
        return "\n".join(lines) if lines else "This presentation has no slides."


class ListPresentationsTool(Tool):
    name = "list_presentations"
    description = "List all presentations Jarvis has created."
    input_schema = {"type": "object", "properties": {}}

    def __init__(self, store: Store):
        self._store = store

    def run(self) -> str:
        docs = self._store.list_documents(kind="pptx")
        if not docs:
            return "No presentations yet."
        return "\n".join(f"- {d.name} ({d.path})" for d in docs)
